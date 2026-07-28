"""
Agent AXE E-SANTE.

Contrairement aux agents précédents (ADOPT/AEMSOFT/AMPLIFON, basés sur
texte_mail seul), AXE E-SANTE.docx précise que la SOURCE PRINCIPALE des
données est une pièce jointe Excel ("demande.xlsx") — le mail lui-même
n'est qu'un mot d'accompagnement, parfois porteur d'infos complémentaires
(confirmé par Sofiane : "il faut toujours checker mail + fichier xlsx").

Une 2ᵉ pièce jointe ("fiche d'intervention" en PDF) n'a PAS besoin d'être
lue : elle part telle quelle dans Documents (sera signée plus tard par le
technicien sur site).

Architecture (décision prise avec Sofiane) : lecture du xlsx 100%
déterministe (openpyxl), PAS via Gemini — c'est une donnée structurée à
libellés connus, un LLM n'apporte rien pour localiser une cellule.

Structure réelle observée (fichier "Matrice Inter IRIS" fourni en exemple) :
- Une seule feuille utile, colonne B uniquement.
- Alternance label / valeur, MAIS avec un titre de section sans valeur
  propre ("Site d'intervention") qui casse l'alternance stricte par numéro
  de ligne -> lecture par CORRESPONDANCE DE LIBELLÉ (`lire_champs_excel`),
  pas par position fixe, pour rester robuste à ce genre de variation.
- Le bloc "Adresse" contient en réalité 3 informations en texte libre
  multi-lignes : "Nom du site : <enseigne>", "<rue>, <CP> <Ville>", et
  "Contact sur site pour intervention : <contact>" -> parsées ensemble par
  `parser_bloc_adresse`.
- Les labels séparés "Contact" / "Téléphone" (plus bas dans le formulaire)
  ne sont PAS le contact site : ils contiennent les coordonnées d'AXE
  E-Santé elle-même (l'intermédiaire) -> volontairement non mappés sur
  customer.prenom/nom/téléphone (cf. hypothèse 3).

⚠️ Hypothèses à vérifier :

  1. Le contrat unique AXE E-SANTE n'a pas de libellé Pivot documenté
     (comme pour AEMSOFT) -> CONTRAT_AXE_ESANTE est laissé vide, à compléter.
  2. "Besoin de matériel (oui/non)" est en pratique du texte libre, pas un
     vrai booléen -> interprété par heuristique, TOUJOURS signalé en
     commentaire, jamais appliqué en silence.
  3. "Contact" / "Téléphone" (labels isolés) = coordonnées d'AXE E-Santé,
     pas du contact site -> non mappés sur customer.*.
  4. "Type de demande" ne contient pas toujours littéralement "Installation"
     ou "Remplacement" -> `detecter_type_demande` croise avec "Description
     complète" en repli.
  5. "Retour de pièces" est déduit du champ "accessoires à récupérer" par
     heuristique.
  6. Pas de champ "Code Site" pour AXE E-SANTE.

--- Intégration Rule Engine (cette étape) ---
  7. Ajout de 2 paramètres optionnels à `enrich_ticket_depuis_excel`, en
     FIN de signature (après `texte_mail`, pour préserver la compatibilité
     positionnelle des appels existants) : `rag_decision: RagDecision |
     None = None` et `activer_rule_engine: bool = False`. Désactivé par
     défaut : rétrocompatibilité totale. AXE E-SANTE n'a qu'UN SEUL point
     de sortie.
"""

import re
import unicodedata

from openpyxl import load_workbook

from app.models.ticket import Ticket
from app.models.rag_decision import RagDecision
from app.services import rule_engine


def _sans_accents(texte: str) -> str:
    """Retire les accents pour fiabiliser les comparaisons de texte."""
    if not texte:
        return ""
    return "".join(
        c for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )


def _normaliser(texte: str) -> str:
    """Normalise (minuscule + sans accent + trim) pour des comparaisons robustes."""
    return _sans_accents((texte or "").strip().lower())


def _ajouter_si_absent(texte_existant: str, bloc: str) -> str:
    """Ajoute `bloc` à `texte_existant` s'il n'y est pas déjà (idempotence)."""
    if not bloc:
        return texte_existant
    if texte_existant and bloc in texte_existant:
        return texte_existant
    if texte_existant:
        return f"{texte_existant.strip()}\n\n{bloc}"
    return bloc


def _formater_recommandations_rule_engine(recommandations) -> str:
    """
    Formate les RuleRecommendation (rule_engine.executer) en un bloc de
    texte destiné à commentaire_interne -- ne modifie JAMAIS un champ
    métier directement (même politique que sur les autres agents migrés).
    """
    if not recommandations:
        return ""
    lignes = ["🧩 Recommandations du Rule Engine (à vérifier, jamais appliquées automatiquement) :"]
    for reco in recommandations:
        lignes.append(
            f"- Champ '{reco.field}' -> '{reco.value}' "
            f"(confiance={reco.confidence:.2f}, source={reco.source}) : {reco.reason}"
        )
    return "\n".join(lignes)


NUMERO_CLIENT_AXE_ESANTE = "DIFFME21"
CONTRAT_AXE_ESANTE = ""  # TODO : libellé Pivot non documenté, cf. hypothèse 1

LABELS_A_MAPPER = {
    "intitule": ("intitule de la demande", "startswith"),
    "type_demande": ("type de demande", "startswith"),
    "adresse_bloc": ("adresse", "exact"),
    "date_heure": ("date et heure intervention souhaitee", "startswith"),
    "description": ("description complete de la demande", "startswith"),
    "temps_estime": ("temps estime sur place", "startswith"),
    "besoin_materiel": ("besoin de materiel", "startswith"),
    "accessoires": ("y a-t-il des accessoires", "startswith"),
    "outillage": ("materiel specifique necessaire", "startswith"),
    "commentaire": ("commentaire si besoin", "startswith"),
    "documents": ("documents a joindre", "startswith"),
    "_entite_facturation": ("entite de facturation", "startswith"),
    "_contact_axe": ("contact", "exact"),
    "_telephone_axe": ("telephone", "exact"),
}

RE_DUREE_H = re.compile(r"(\d+(?:[.,]\d+)?)\s*h", re.IGNORECASE)
RE_NB_TECH = re.compile(r"(\d+)\s*tech", re.IGNORECASE)
RE_NOM_SITE = re.compile(r"nom du site\s*:\s*(.+)", re.IGNORECASE)
RE_CONTACT_SITE = re.compile(r"contact sur site pour intervention\s*:\s*(.+)", re.IGNORECASE)
RE_RUE_CP_VILLE = re.compile(r"(.+?),\s*(\d{5})\s+(.+)")


def lire_champs_excel(fichier, feuille: str | None = None) -> dict:
    """
    Lit le xlsx AXE E-SANTE (formulaire vertical colonne B, label puis
    valeur en alternance -- avec quelques lignes de bruit).
    """
    wb = load_workbook(fichier, data_only=True)
    ws = wb[feuille] if feuille else wb.active

    cellules = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                cellules.append(str(cell.value).strip())

    resultat: dict[str, str] = {}
    i = 0
    while i < len(cellules):
        texte_normalise = _normaliser(cellules[i])
        cle_trouvee = None
        for cle, (motif, mode) in LABELS_A_MAPPER.items():
            if mode == "exact" and texte_normalise == motif:
                cle_trouvee = cle
                break
            if mode == "startswith" and texte_normalise.startswith(motif):
                cle_trouvee = cle
                break
        if cle_trouvee and i + 1 < len(cellules):
            if not cle_trouvee.startswith("_"):
                resultat[cle_trouvee] = cellules[i + 1]
            i += 2
        else:
            i += 1
    return resultat


def parser_bloc_adresse(bloc: str) -> dict:
    """Parse le bloc "Adresse" (texte libre multi-lignes)."""
    resultat = {"enseigne": "", "adresse": "", "code_postal": "", "ville": "", "contact_site": ""}
    if not bloc:
        return resultat
    for ligne in (l.strip() for l in bloc.splitlines() if l.strip()):
        m_site = RE_NOM_SITE.match(ligne)
        if m_site:
            resultat["enseigne"] = m_site.group(1).strip()
            continue
        m_contact = RE_CONTACT_SITE.match(ligne)
        if m_contact:
            resultat["contact_site"] = m_contact.group(1).strip()
            continue
        m_adresse = RE_RUE_CP_VILLE.match(ligne)
        if m_adresse:
            resultat["adresse"] = m_adresse.group(1).strip()
            resultat["code_postal"] = m_adresse.group(2).strip()
            resultat["ville"] = m_adresse.group(3).strip()
    return resultat


def detecter_type_demande(type_brut: str, description: str) -> str:
    """Installation / Remplacement."""
    texte = _normaliser(f"{type_brut or ''} {description or ''}")
    if "sans integration" in texte:
        return "Installation"
    if "remplacement" in texte:
        return "Remplacement"
    if "installation" in texte or "integration" in texte:
        return "Installation"
    return ""


def extraire_duree_et_techniciens(valeur: str) -> tuple[str, int]:
    """Parse '1h 1tech' -> ('1h', 1)."""
    duree, nb_tech = "", 0
    if not valeur:
        return duree, nb_tech
    m_duree = RE_DUREE_H.search(valeur)
    if m_duree:
        duree = f"{m_duree.group(1)}h"
    m_tech = RE_NB_TECH.search(valeur)
    if m_tech:
        nb_tech = int(m_tech.group(1))
    return duree, nb_tech


def interpreter_besoin_materiel(valeur: str) -> tuple[bool, str]:
    """Le label demande oui/non mais la valeur réelle est souvent une phrase libre."""
    texte = _normaliser(valeur)
    if not texte:
        return False, "Besoin de matériel non renseigné dans le xlsx — laissé à Non par défaut, à vérifier."
    if "disponible chez le client" in texte or "deja sur place" in texte or "deja disponible" in texte or re.search(r"\bnon\b", texte):
        return False, f"Besoin de matériel interprété comme Non depuis le texte libre du xlsx : {valeur!r}."
    if texte == "oui" or texte.startswith("oui"):
        return True, f"Besoin de matériel interprété comme Oui depuis le texte libre du xlsx : {valeur!r}."
    return False, (
        f"Besoin de matériel AMBIGU dans le xlsx (valeur brute : {valeur!r}) "
        f"— interprétation par défaut = Non, À VÉRIFIER MANUELLEMENT."
    )


def interpreter_retour_piece(valeur: str) -> tuple[str, str]:
    """Retour de pièces déduit du texte libre 'accessoires à récupérer'."""
    texte = _normaliser(valeur)
    if not texte:
        return "", ""
    if "laisser" in texte and ("place" in texte or "client" in texte):
        return "Non", f"Retour de pièces interprété comme Non : {valeur!r}."
    if "recuperer" in texte or "reprendre" in texte or "retour" in texte:
        return "Oui", f"Retour de pièces interprété comme Oui : {valeur!r}."
    return "", f"Retour de pièces ambigu (valeur brute : {valeur!r}) — à vérifier manuellement."


def construire_intitule(type_demande: str, intitule_brut: str) -> str:
    """Intitulé = '<Type> – <Intitulé de la demande>'."""
    intitule_brut = (intitule_brut or "").strip()
    if not type_demande or not intitule_brut:
        return intitule_brut
    return f"{type_demande} – {intitule_brut}"


def construire_travail_attendu(champs: dict) -> str:
    """Travail attendu = Description + Besoin de matériel + Accessoires + Commentaire + Documents."""
    parties = []
    if champs.get("description"):
        parties.append(champs["description"])
    if champs.get("besoin_materiel"):
        parties.append(f"Besoin de matériel (texte original) : {champs['besoin_materiel']}")
    if champs.get("accessoires"):
        parties.append(f"Accessoires à récupérer : {champs['accessoires']}")
    if champs.get("commentaire"):
        parties.append(f"Commentaire : {champs['commentaire']}")
    if champs.get("documents"):
        parties.append(f"Documents à joindre / actions sur site : {champs['documents']}")
    return "\n\n".join(parties)


def enrich_ticket_depuis_excel(
    ticket: Ticket,
    fichier_excel,
    texte_mail: str = "",
    rag_decision: RagDecision | None = None,
    activer_rule_engine: bool = False,
) -> Ticket:
    """
    Enrichit un Ticket à partir du xlsx AXE E-SANTE (source principale) et,
    en complément, du texte du mail.

    `fichier_excel` : chemin ou objet fichier du "demande.xlsx". Ne PAS
    passer le PDF "fiche d'intervention" ici.

    `rag_decision` (optionnel) : une RagDecision déjà calculée en amont,
    transmise telle quelle au Rule Engine si celui-ci est activé.

    `activer_rule_engine` (par défaut False) : si True, exécute
    rule_engine.executer(ticket, rag_decision) et ajoute ses
    recommandations à commentaire_interne -- jamais à un champ métier.
    """
    notes: list[str] = []

    champs = lire_champs_excel(fichier_excel)

    ticket.customer.client = "AXE E-SANTE"
    ticket.customer.numero_client = NUMERO_CLIENT_AXE_ESANTE
    ticket.customer.pays = "France"

    bloc_adresse = parser_bloc_adresse(champs.get("adresse_bloc", ""))
    if bloc_adresse["enseigne"]:
        ticket.customer.enseigne = bloc_adresse["enseigne"]
    if bloc_adresse["adresse"]:
        ticket.customer.adresse = bloc_adresse["adresse"]
    if bloc_adresse["code_postal"]:
        ticket.customer.code_postal = bloc_adresse["code_postal"]
    if bloc_adresse["ville"]:
        ticket.customer.ville = bloc_adresse["ville"]
    if bloc_adresse["contact_site"]:
        ticket.customer.nom = bloc_adresse["contact_site"]
    else:
        notes.append("Contact sur site introuvable dans le bloc 'Adresse' du xlsx — à compléter manuellement.")

    ticket.intervention.type_intervention = "Contrat"
    if CONTRAT_AXE_ESANTE:
        ticket.intervention.contrat = CONTRAT_AXE_ESANTE
    elif not ticket.intervention.contrat:
        notes.append(
            "Contrat AXE E-SANTE non renseigné automatiquement : un seul contrat "
            "existe côté Pivot, mais son libellé exact n'est pas documenté dans "
            "AXE_E-SANTE.docx — sélectionner manuellement, ou compléter la "
            "constante CONTRAT_AXE_ESANTE dans cet agent."
        )
    else:
        notes.append(
            f"Contrat AXE E-SANTE déjà renseigné ({ticket.intervention.contrat!r}) "
            f"mais NON VÉRIFIABLE : son libellé Pivot exact n'est documenté nulle "
            f"part — à confirmer avant saisie."
        )

    type_demande = detecter_type_demande(champs.get("type_demande", ""), champs.get("description", ""))
    if not type_demande:
        notes.append(
            f"Type de demande non reconnu comme Installation/Remplacement "
            f"(valeur brute : {champs.get('type_demande', '')!r}) — à choisir manuellement."
        )
    ticket.intervention.type = type_demande

    ticket.intervention.intitule = construire_intitule(type_demande, champs.get("intitule", ""))
    ticket.intervention.problematique = champs.get("description", "")
    ticket.intervention.origine = "Email"

    besoin_materiel, note_besoin = interpreter_besoin_materiel(champs.get("besoin_materiel", ""))
    ticket.logistics.besoin_materiel = besoin_materiel
    if note_besoin:
        notes.append(note_besoin)

    retour_piece, note_retour = interpreter_retour_piece(champs.get("accessoires", ""))
    if retour_piece:
        ticket.logistics.retour_piece = retour_piece
    if note_retour:
        notes.append(note_retour)

    if ticket.logistics.besoin_materiel:
        ticket.logistics.envoi_piece_par = "IRIS"
        ticket.logistics.consigne_livraison = "Pudo"
        notes.append("Besoin de matériel = Oui : vérifier si le client demande une livraison directe plutôt que Pudo.")

    duree, nb_tech = extraire_duree_et_techniciens(champs.get("temps_estime", ""))
    if duree:
        ticket.procedure.duree = duree
    if nb_tech:
        ticket.procedure.nombre_techniciens = nb_tech
    if not duree or not nb_tech:
        notes.append(
            f"Durée/Nombre de techniciens partiellement extrait de 'Temps estimé sur "
            f"place' (valeur brute : {champs.get('temps_estime', '')!r}) — à vérifier."
        )

    if champs.get("date_heure"):
        ticket.procedure.date_limite = champs["date_heure"]
        ticket.procedure.contrainte = f"Date/heure imposée par le client : {champs['date_heure']}"

    ticket.procedure.travail_attendu = construire_travail_attendu(champs)

    if champs.get("outillage"):
        ticket.procedure.autre_outillage = champs["outillage"]

    ticket.procedure.intervention_sur_site = True
    ticket.procedure.technicien_anglophone = False
    ticket.procedure.procedure = True

    if texte_mail and texte_mail.strip():
        notes.append(
            "Le mail contenait du texte en plus du xlsx — à relire pour d'éventuelles "
            "informations complémentaires (non extraites automatiquement par cet agent)."
        )

    if notes:
        bloc_notes = "⚠️ Points à vérifier (générés automatiquement) :\n" + "\n".join(f"- {n}" for n in notes)
        ticket.intervention.commentaire_interne = _ajouter_si_absent(
            ticket.intervention.commentaire_interne, bloc_notes
        )

    if activer_rule_engine:
        recommandations = rule_engine.executer(ticket, rag_decision)
        bloc_recommandations = _formater_recommandations_rule_engine(recommandations)
        ticket.intervention.commentaire_interne = _ajouter_si_absent(
            ticket.intervention.commentaire_interne, bloc_recommandations
        )

    return ticket

"""
Agent AXE E-SANTE.

[... docstring métier inchangé ...]

--- Harmonisation Framework des Agents (P3-424) ---
Logique métier déplacée dans `_appliquer_regles_axe_esante` (privée),
appelée par `AxeEsanteAgent.analyze()`. `analyze()` ne pilote jamais le
Rule Engine (Règle 9). `enrich_ticket_depuis_excel()` reste un adaptateur
de compatibilité.

Contrairement à PROMETHEAN, ce client s'intègre NATIVEMENT au contrat
`AgentRequest` gelé : `request.fichier_attache` porte directement le
xlsx, aucun compromis nécessaire.
"""

import re
import unicodedata

from openpyxl import load_workbook

from app.agents.base_agent import BaseAgent
from app.models.agent_contracts import AgentRequest, AgentResult
from app.models.rag_decision import RagDecision
from app.models.ticket import Ticket
from app.services import rule_engine


def _sans_accents(texte: str) -> str:
    if not texte:
        return ""
    return "".join(
        c for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )


def _normaliser(texte: str) -> str:
    return _sans_accents((texte or "").strip().lower())


def _ajouter_si_absent(texte_existant: str, bloc: str) -> str:
    if not bloc:
        return texte_existant
    if texte_existant and bloc in texte_existant:
        return texte_existant
    if texte_existant:
        return f"{texte_existant.strip()}\n\n{bloc}"
    return bloc


def _formater_recommandations_rule_engine(recommandations) -> str:
    if not recommandations:
        return ""
    lignes = ["🧩 Recommandations du Rule Engine (à vérifier, jamais appliquées automatiquement) :"]
    for reco in recommandations:
        lignes.append(
            f"- Champ '{reco.field}' -> '{reco.value}' "
            f"(confiance={reco.confidence:.2f}, source={reco.source}) : {reco.reason}"
        )
    return "\n".join(lignes)


NUMERO_CLIENT_AXE_ESANTE = "DIFFME21"
CONTRAT_AXE_ESANTE = ""

LABELS_A_MAPPER = {
    "intitule": ("intitule de la demande", "startswith"),
    "type_demande": ("type de demande", "startswith"),
    "adresse_bloc": ("adresse", "exact"),
    "date_heure": ("date et heure intervention souhaitee", "startswith"),
    "description": ("description complete de la demande", "startswith"),
    "temps_estime": ("temps estime sur place", "startswith"),
    "besoin_materiel": ("besoin de materiel", "startswith"),
    "accessoires": ("y a-t-il des accessoires", "startswith"),
    "outillage": ("materiel specifique necessaire", "startswith"),
    "commentaire": ("commentaire si besoin", "startswith"),
    "documents": ("documents a joindre", "startswith"),
    "_entite_facturation": ("entite de facturation", "startswith"),
    "_contact_axe": ("contact", "exact"),
    "_telephone_axe": ("telephone", "exact"),
}

RE_DUREE_H = re.compile(r"(\d+(?:[.,]\d+)?)\s*h", re.IGNORECASE)
RE_NB_TECH = re.compile(r"(\d+)\s*tech", re.IGNORECASE)
RE_NOM_SITE = re.compile(r"nom du site\s*:\s*(.+)", re.IGNORECASE)
RE_CONTACT_SITE = re.compile(r"contact sur site pour intervention\s*:\s*(.+)", re.IGNORECASE)
RE_RUE_CP_VILLE = re.compile(r"(.+?),\s*(\d{5})\s+(.+)")


def lire_champs_excel(fichier, feuille: str | None = None) -> dict:
    wb = load_workbook(fichier, data_only=True)
    ws = wb[feuille] if feuille else wb.active

    cellules = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                cellules.append(str(cell.value).strip())

    resultat: dict[str, str] = {}
    i = 0
    while i < len(cellules):
        texte_normalise = _normaliser(cellules[i])
        cle_trouvee = None
        for cle, (motif, mode) in LABELS_A_MAPPER.items():
            if mode == "exact" and texte_normalise == motif:
                cle_trouvee = cle
                break
            if mode == "startswith" and texte_normalise.startswith(motif):
                cle_trouvee = cle
                break
        if cle_trouvee and i + 1 < len(cellules):
            if not cle_trouvee.startswith("_"):
                resultat[cle_trouvee] = cellules[i + 1]
            i += 2
        else:
            i += 1
    return resultat


def parser_bloc_adresse(bloc: str) -> dict:
    resultat = {"enseigne": "", "adresse": "", "code_postal": "", "ville": "", "contact_site": ""}
    if not bloc:
        return resultat
    for ligne in (l.strip() for l in bloc.splitlines() if l.strip()):
        m_site = RE_NOM_SITE.match(ligne)
        if m_site:
            resultat["enseigne"] = m_site.group(1).strip()
            continue
        m_contact = RE_CONTACT_SITE.match(ligne)
        if m_contact:
            resultat["contact_site"] = m_contact.group(1).strip()
            continue
        m_adresse = RE_RUE_CP_VILLE.match(ligne)
        if m_adresse:
            resultat["adresse"] = m_adresse.group(1).strip()
            resultat["code_postal"] = m_adresse.group(2).strip()
            resultat["ville"] = m_adresse.group(3).strip()
    return resultat


def detecter_type_demande(type_brut: str, description: str) -> str:
    texte = _normaliser(f"{type_brut or ''} {description or ''}")
    if "sans integration" in texte:
        return "Installation"
    if "remplacement" in texte:
        return "Remplacement"
    if "installation" in texte or "integration" in texte:
        return "Installation"
    return ""


def extraire_duree_et_techniciens(valeur: str) -> tuple[str, int]:
    duree, nb_tech = "", 0
    if not valeur:
        return duree, nb_tech
    m_duree = RE_DUREE_H.search(valeur)
    if m_duree:
        duree = f"{m_duree.group(1)}h"
    m_tech = RE_NB_TECH.search(valeur)
    if m_tech:
        nb_tech = int(m_tech.group(1))
    return duree, nb_tech


def interpreter_besoin_materiel(valeur: str) -> tuple[bool, str]:
    texte = _normaliser(valeur)
    if not texte:
        return False, "Besoin de matériel non renseigné dans le xlsx — laissé à Non par défaut, à vérifier."
    if "disponible chez le client" in texte or "deja sur place" in texte or "deja disponible" in texte or re.search(r"\bnon\b", texte):
        return False, f"Besoin de matériel interprété comme Non depuis le texte libre du xlsx : {valeur!r}."
    if texte == "oui" or texte.startswith("oui"):
        return True, f"Besoin de matériel interprété comme Oui depuis le texte libre du xlsx : {valeur!r}."
    return False, (
        f"Besoin de matériel AMBIGU dans le xlsx (valeur brute : {valeur!r}) "
        f"— interprétation par défaut = Non, À VÉRIFIER MANUELLEMENT."
    )


def interpreter_retour_piece(valeur: str) -> tuple[str, str]:
    texte = _normaliser(valeur)
    if not texte:
        return "", ""
    if "laisser" in texte and ("place" in texte or "client" in texte):
        return "Non", f"Retour de pièces interprété comme Non : {valeur!r}."
    if "recuperer" in texte or "reprendre" in texte or "retour" in texte:
        return "Oui", f"Retour de pièces interprété comme Oui : {valeur!r}."
    return "", f"Retour de pièces ambigu (valeur brute : {valeur!r}) — à vérifier manuellement."


def construire_intitule(type_demande: str, intitule_brut: str) -> str:
    intitule_brut = (intitule_brut or "").strip()
    if not type_demande or not intitule_brut:
        return intitule_brut
    return f"{type_demande} – {intitule_brut}"


def construire_travail_attendu(champs: dict) -> str:
    parties = []
    if champs.get("description"):
        parties.append(champs["description"])
    if champs.get("besoin_materiel"):
        parties.append(f"Besoin de matériel (texte original) : {champs['besoin_materiel']}")
    if champs.get("accessoires"):
        parties.append(f"Accessoires à récupérer : {champs['accessoires']}")
    if champs.get("commentaire"):
        parties.append(f"Commentaire : {champs['commentaire']}")
    if champs.get("documents"):
        parties.append(f"Documents à joindre / actions sur site : {champs['documents']}")
    return "\n\n".join(parties)


def _appliquer_regles_axe_esante(ticket: Ticket, fichier_excel, texte_mail: str = "") -> Ticket:
    """Logique métier AXE E-SANTE pure (extraite de l'ancien `enrich_ticket_depuis_excel`, sans le bloc Rule Engine)."""
    notes: list[str] = []

    champs = lire_champs_excel(fichier_excel)

    ticket.customer.client = "AXE E-SANTE"
    ticket.customer.numero_client = NUMERO_CLIENT_AXE_ESANTE
    ticket.customer.pays = "France"

    bloc_adresse = parser_bloc_adresse(champs.get("adresse_bloc", ""))
    if bloc_adresse["enseigne"]:
        ticket.customer.enseigne = bloc_adresse["enseigne"]
    if bloc_adresse["adresse"]:
        ticket.customer.adresse = bloc_adresse["adresse"]
    if bloc_adresse["code_postal"]:
        ticket.customer.code_postal = bloc_adresse["code_postal"]
    if bloc_adresse["ville"]:
        ticket.customer.ville = bloc_adresse["ville"]
    if bloc_adresse["contact_site"]:
        ticket.customer.nom = bloc_adresse["contact_site"]
    else:
        notes.append("Contact sur site introuvable dans le bloc 'Adresse' du xlsx — à compléter manuellement.")

    ticket.intervention.type_intervention = "Contrat"
    if CONTRAT_AXE_ESANTE:
        ticket.intervention.contrat = CONTRAT_AXE_ESANTE
    elif not ticket.intervention.contrat:
        notes.append(
            "Contrat AXE E-SANTE non renseigné automatiquement : un seul contrat "
            "existe côté Pivot, mais son libellé exact n'est pas documenté dans "
            "AXE_E-SANTE.docx — sélectionner manuellement, ou compléter la "
            "constante CONTRAT_AXE_ESANTE dans cet agent."
        )
    else:
        notes.append(
            f"Contrat AXE E-SANTE déjà renseigné ({ticket.intervention.contrat!r}) "
            f"mais NON VÉRIFIABLE : son libellé Pivot exact n'est documenté nulle "
            f"part — à confirmer avant saisie."
        )

    type_demande = detecter_type_demande(champs.get("type_demande", ""), champs.get("description", ""))
    if not type_demande:
        notes.append(
            f"Type de demande non reconnu comme Installation/Remplacement "
            f"(valeur brute : {champs.get('type_demande', '')!r}) — à choisir manuellement."
        )
    ticket.intervention.type = type_demande

    ticket.intervention.intitule = construire_intitule(type_demande, champs.get("intitule", ""))
    ticket.intervention.problematique = champs.get("description", "")
    ticket.intervention.origine = "Email"

    besoin_materiel, note_besoin = interpreter_besoin_materiel(champs.get("besoin_materiel", ""))
    ticket.logistics.besoin_materiel = besoin_materiel
    if note_besoin:
        notes.append(note_besoin)

    retour_piece, note_retour = interpreter_retour_piece(champs.get("accessoires", ""))
    if retour_piece:
        ticket.logistics.retour_piece = retour_piece
    if note_retour:
        notes.append(note_retour)

    if ticket.logistics.besoin_materiel:
        ticket.logistics.envoi_piece_par = "IRIS"
        ticket.logistics.consigne_livraison = "Pudo"
        notes.append("Besoin de matériel = Oui : vérifier si le client demande une livraison directe plutôt que Pudo.")

    duree, nb_tech = extraire_duree_et_techniciens(champs.get("temps_estime", ""))
    if duree:
        ticket.procedure.duree = duree
    if nb_tech:
        ticket.procedure.nombre_techniciens = nb_tech
    if not duree or not nb_tech:
        notes.append(
            f"Durée/Nombre de techniciens partiellement extrait de 'Temps estimé sur "
            f"place' (valeur brute : {champs.get('temps_estime', '')!r}) — à vérifier."
        )

    if champs.get("date_heure"):
        ticket.procedure.date_limite = champs["date_heure"]
        ticket.procedure.contrainte = f"Date/heure imposée par le client : {champs['date_heure']}"

    ticket.procedure.travail_attendu = construire_travail_attendu(champs)

    if champs.get("outillage"):
        ticket.procedure.autre_outillage = champs["outillage"]

    ticket.procedure.intervention_sur_site = True
    ticket.procedure.technicien_anglophone = False
    ticket.procedure.procedure = True

    if texte_mail and texte_mail.strip():
        notes.append(
            "Le mail contenait du texte en plus du xlsx — à relire pour d'éventuelles "
            "informations complémentaires (non extraites automatiquement par cet agent)."
        )

    if notes:
        bloc_notes = "⚠️ Points à vérifier (générés automatiquement) :\n" + "\n".join(f"- {n}" for n in notes)
        ticket.intervention.commentaire_interne = _ajouter_si_absent(
            ticket.intervention.commentaire_interne, bloc_notes
        )

    return ticket


class AxeEsanteAgent(BaseAgent):
    """Agent AXE E-SANTE conforme au contrat BaseAgent. `request.fichier_attache` porte le xlsx directement."""

    def analyze(self, request: AgentRequest) -> AgentResult:
        try:
            ticket = _appliquer_regles_axe_esante(
                request.ticket, request.fichier_attache, texte_mail=request.texte_mail
            )
            return AgentResult(ticket=ticket, succes=True)
        except Exception as exc:
            return AgentResult(ticket=request.ticket, succes=False, erreur=str(exc))


_AGENT = AxeEsanteAgent()


def enrich_ticket_depuis_excel(
    ticket: Ticket,
    fichier_excel,
    texte_mail: str = "",
    rag_decision: RagDecision | None = None,
    activer_rule_engine: bool = False,
) -> Ticket:
    """⚠️ ADAPTATEUR DE COMPATIBILITÉ -- délègue à AxeEsanteAgent.analyze(), reproduit ici l'ancien comportement de activer_rule_engine."""
    request = AgentRequest(ticket=ticket, texte_mail=texte_mail, fichier_attache=fichier_excel, rag_decision=rag_decision)
    result = _AGENT.analyze(request)
    ticket_resultat = result.ticket

    if activer_rule_engine:
        recommandations = rule_engine.executer(ticket_resultat, rag_decision)
        bloc_recommandations = _formater_recommandations_rule_engine(recommandations)
        ticket_resultat.intervention.commentaire_interne = _ajouter_si_absent(
            ticket_resultat.intervention.commentaire_interne, bloc_recommandations
        )
    return ticket_resultat