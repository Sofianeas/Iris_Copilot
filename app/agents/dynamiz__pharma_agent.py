"""
Agent DYNAMIZ PHARMA.

Logique métier issue de DYNAMIZ_PHARMA.docx + SKILL.md.

⚠️ Contrairement à AXE E-SANTE (vrai fichier inspecté) et ETAM (capture
d'écran), je n'ai ici qu'une RETRANSCRIPTION TEXTE d'un exemple -- pas de
fichier réel. Sofiane précise aussi que les informations sont parfois dans
le mail, parfois dans la pièce jointe Excel, parfois les deux -> plutôt que
de deviner une mise en page de cellules (fragile), cet agent travaille sur
du TEXTE BRUT par extraction regex, indépendamment de la source exacte
(mail, Excel aplati en texte, ou les deux concaténés). C'est un choix
architectural différent d'AXE E-SANTE/ETAM, justifié par cette variabilité
de source plutôt que par une mise en page de cellules fixe.

Règles clés (DYNAMIZ_PHARMA.docx) :
- Pays = toujours France. Technicien anglophone = toujours Non.
- ⚠️ Procédure = NON par défaut chez ce client -- c'est le SEUL client de
  toute la base où c'est le cas (tous les autres ont Oui). Erreur courante
  à éviter, cf. SKILL.md "erreurs courantes".
- Code Site = le "code client" du bloc "Client à facturer" (PAS le champ
  "Code site" séparé sous "Site d'intervention", souvent vide).
- "Description de la demande" se décompose en 2 sous-sections : "Contexte:"
  -> Problématique, "Intervention à réaliser:" (+ tout ce qui suit) ->
  Travail attendu.
- Type : INSTALLATION / MAINTENANCE SAV / PREVISITE (champ "Intitulé de la
  demande" du formulaire, valeurs SAV/INSTALLATION/PREVISITE).
- Type de ticket : Incident si SAV (Maintenance), Demande si
  INSTALLATION/PREVISITE.
- "Hauteur de l'intervention" -> Autre outillage spécifique. Si le texte
  mentionne la nécessité de confirmer qu'un câble réseau est bien actif/
  connecté -> ajouter "Testeur de câbles" (sans inclure le matériel déjà
  envoyé).
- "Date et heure intervention souhaitée" : si ce n'est pas une vraie date/
  heure mais une consigne de prise de contact (cas fréquent, cf. exemple
  fourni), ce n'est PAS une Date limite -> classé en Consigne de
  planification à la place.

⚠️ Hypothèses à vérifier :
  1. Mise en page exacte non vérifiée (cf. avertissement ci-dessus) --
     l'extraction repose sur la présence des libellés exacts observés dans
     l'exemple fourni, pas sur une structure de cellules confirmée.
  2. Contrat unique non documenté (même situation qu'AEMSOFT/AXE E-SANTE)
     -> CONTRAT_DYNAMIZ laissé vide, à compléter une fois confirmé.
  3. "Intitulé" construit selon le seul exemple donné dans le docx
     ("Maintenance Player ..." / "PréVisite ...") -- le préfixe "Player"
     est peut-être spécifique à ce type d'équipement (écran Cenareo), pas
     une règle générale pour tout type de matériel SAV. À confirmer avec
     un 2ᵉ exemple impliquant un autre type de matériel.
"""

import re
import unicodedata

from app.models.ticket import Ticket


def _sans_accents(texte: str) -> str:
    """Retire les accents pour fiabiliser les comparaisons de texte."""
    if not texte:
        return ""
    return "".join(
        c for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )


def _normaliser(texte: str) -> str:
    """Normalise (minuscule + sans accent + trim) pour des comparaisons robustes."""
    return _sans_accents((texte or "").strip().lower())


def _ajouter_si_absent(texte_existant: str, bloc: str) -> str:
    """Ajoute `bloc` à `texte_existant` s'il n'y est pas déjà (idempotence)."""
    if not bloc:
        return texte_existant
    if texte_existant and bloc in texte_existant:
        return texte_existant
    if texte_existant:
        return f"{texte_existant.strip()}\n\n{bloc}"
    return bloc


# --------------------------------------------------------------------------
# Lecture d'un fichier Excel -> texte brut (réutilise les mêmes extracteurs
# regex que pour le texte du mail, cf. avertissement du docstring)
# --------------------------------------------------------------------------

def lire_excel_vers_texte(fichier, feuille: str | None = None) -> str:
    """Aplati toutes les cellules d'un xlsx en texte (1 valeur par ligne)."""
    from openpyxl import load_workbook
    wb = load_workbook(fichier, data_only=True)
    ws = wb[feuille] if feuille else wb.active
    lignes = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                lignes.append(str(cell.value))
    return "\n".join(lignes)


# --------------------------------------------------------------------------
# Référentiel DYNAMIZ PHARMA
# --------------------------------------------------------------------------

CONTRAT_DYNAMIZ = ""  # TODO : libellé Pivot non documenté, cf. hypothèse 2

TYPE_INSTALLATION = "INSTALLATION"
TYPE_MAINTENANCE_SAV = "MAINTENANCE SAV"
TYPE_PREVISITE = "PREVISITE"

# Séquence ordonnée des libellés du formulaire : chaque champ est borné par
# le libellé suivant dans cette liste (extraction "tout ce qu'il y a entre
# ce libellé et le prochain").
SEQUENCE_LABELS = [
    ("client_facturation", r"client\s+[àa]\s+facturer"),
    ("type_demande", r"intitul[ée] de la demande"),
    ("description", r"description de la demande"),
    ("numero_incident_client", r"num[ée]ro incident client"),
    ("code_site_champ", r"\bcode site\b"),
    ("adresse", r"\badresse\b"),
    ("cp_ville", r"\bcp\s*:"),
    ("nom_point_de_vente", r"nom point de vente"),
    ("contact_site", r"contact sur site"),
    ("telephone", r"t[ée]l[ée]phone"),
    ("date_heure", r"date et heure intervention souhait[ée]e"),
    ("hauteur", r"hauteur de l'intervention"),
    ("contraintes", r"contraintes de mises? en place"),
    ("besoin_materiel", r"besoin de mat[ée]riel"),
    ("reference_materiel", r"si oui\s*,?\s*r[ée]f[ée]rence\s+[àa]\s+prendre"),
    ("materiel_a_recuperer", r"mat[ée]riel\s+[àa]\s+r[ée]cup[ée]rer"),
    ("commentaire", r"commentaire\s+si besoin"),
    ("documents", r"documents\s+[àa]\s+joindre"),
    ("nombre_technicien", r"nombre de technicien"),
    ("duree", r"dur[ée]e en heure"),
]

RE_CLIENT_FACTURATION = re.compile(r"nom\s*:\s*(.+?)\s*-\s*code client\s*:\s*(\S+)", re.IGNORECASE)
RE_CP_VILLE_INLINE = re.compile(r"cp\s*:\s*(\d{4,5})\s+ville\s*:\s*(.+)", re.IGNORECASE)
RE_CONTEXTE = re.compile(r"contexte\s*:?\s*\n(.+?)(?=intervention\s+[àa]\s+r[ée]aliser)", re.IGNORECASE | re.DOTALL)
RE_INTERVENTION_A_REALISER = re.compile(r"intervention\s+[àa]\s+r[ée]aliser\s*:?\s*\n(.+)", re.IGNORECASE | re.DOTALL)
RE_DATE_PATTERN = re.compile(r"\b(lundi|mardi|mercredi|jeudi|vendredi|samedi|dimanche)\b.{0,20}\d", re.IGNORECASE)


def extraire_champs_dynamiz(texte: str) -> dict:
    """
    Extrait les champs du formulaire DYNAMIZ PHARMA à partir d'un texte
    brut (mail et/ou Excel aplati, cf. avertissement du docstring du module).
    Chaque champ = texte entre son libellé et le libellé suivant de
    SEQUENCE_LABELS.
    """
    if not texte:
        return {}
    resultat: dict[str, str] = {}
    for i, (cle, motif) in enumerate(SEQUENCE_LABELS):
        motif_suivant = SEQUENCE_LABELS[i + 1][1] if i + 1 < len(SEQUENCE_LABELS) else r"\Z"
        pattern = re.compile(rf"{motif}[^\n]*\n?(.*?)(?={motif_suivant})", re.IGNORECASE | re.DOTALL)
        match = pattern.search(texte)
        if match:
            valeur = match.group(1).strip().strip('"').strip()
            resultat[cle] = valeur
    return resultat


def extraire_client_facturation(texte: str) -> tuple[str, str]:
    """'Nom : <nom> - code client : <code>' -> (nom, code_client)."""
    match = RE_CLIENT_FACTURATION.search(texte or "")
    if not match:
        return "", ""
    return match.group(1).strip(), match.group(2).strip()


def extraire_cp_ville(texte: str) -> tuple[str, str]:
    """'CP: <code postal> VILLE: <ville>' -> (code_postal, ville)."""
    match = RE_CP_VILLE_INLINE.search(texte or "")
    if not match:
        return "", ""
    return match.group(1).strip(), match.group(2).strip()


def extraire_contexte_et_intervention(description: str) -> tuple[str, str]:
    """
    Sépare le bloc Description en 2 : 'Contexte:' -> Problématique,
    'Intervention à réaliser:' (+ tout ce qui suit) -> Travail attendu.
    Si ces sous-sections ne sont pas trouvées, retourne (description, "").
    """
    if not description:
        return "", ""
    m_contexte = RE_CONTEXTE.search(description)
    m_intervention = RE_INTERVENTION_A_REALISER.search(description)
    contexte = m_contexte.group(1).strip() if m_contexte else ""
    intervention = m_intervention.group(1).strip() if m_intervention else ""
    if not contexte and not intervention:
        return description.strip(), ""
    return contexte, intervention


def detecter_type_demande(type_brut: str) -> str:
    """SAV -> MAINTENANCE SAV ; sinon INSTALLATION/PREVISITE tels quels."""
    valeur = _normaliser(type_brut)
    if valeur == "sav":
        return TYPE_MAINTENANCE_SAV
    if valeur == "installation":
        return TYPE_INSTALLATION
    if valeur == "previsite":
        return TYPE_PREVISITE
    return ""


def detecter_besoin_testeur_cable(texte: str) -> bool:
    """
    Docx : si le texte mentionne la nécessité de confirmer/vérifier qu'un
    câble réseau est bien actif/connecté -> ajouter "Testeur de câbles"
    dans Autre outillage spécifique.
    """
    texte_n = _normaliser(texte)
    return ("cable" in texte_n) and any(
        mot in texte_n for mot in ("confirmer", "actif", "verifier", "diagnostiquer", "branche")
    )


def ressemble_a_une_date(texte: str) -> bool:
    """Heuristique : un jour de semaine suivi de chiffres (date/heure) à proximité."""
    return bool(RE_DATE_PATTERN.search(texte or ""))


def construire_intitule(type_demande: str, enseigne: str, ville: str) -> str:
    """
    Intitulé = "<Problème/motif> + Enseigne + Ville" (cf. exemples
    DYNAMIZ_PHARMA.docx). "Maintenance Player" / "PréVisite" sont les
    préfixes donnés en exemple dans le docx -- cf. hypothèse 3 (incertain
    pour un matériel autre qu'un player/écran).
    """
    prefixe = {
        TYPE_MAINTENANCE_SAV: "Maintenance Player",
        TYPE_PREVISITE: "PréVisite",
        TYPE_INSTALLATION: "Installation",
    }.get(type_demande, "")
    parts = [p for p in (prefixe, enseigne, ville) if p]
    return " ".join(parts)


# --------------------------------------------------------------------------
# Agent
# --------------------------------------------------------------------------

def enrich_ticket(ticket: Ticket, texte_source: str = "", fichier_excel=None) -> Ticket:
    """
    Enrichit un Ticket à partir d'un texte brut (mail et/ou Excel aplati).

    `texte_source` : texte du mail. `fichier_excel` (optionnel) : chemin/
    objet fichier de la pièce jointe Excel, aplati et concaténé à
    texte_source si fourni (les deux sources sont traitées de façon
    identique, cf. avertissement du docstring du module).
    """
    notes: list[str] = []

    if fichier_excel is not None:
        texte_source = _ajouter_si_absent(texte_source, lire_excel_vers_texte(fichier_excel))

    champs = extraire_champs_dynamiz(texte_source)

    ticket.customer.client = "DYNAMIZ"
    ticket.customer.pays = "France"

    nom_facturation, code_client = extraire_client_facturation(champs.get("client_facturation", ""))
    if code_client:
        ticket.customer.code_site = code_client
        if not re.fullmatch(r"\d{6,8}", code_client):
            notes.append(f"Code client inhabituel ({code_client!r}) — attendu ~7 chiffres, à vérifier.")
    else:
        notes.append("Code client ('Client à facturer' -> code client) introuvable — Code Site à compléter manuellement.")

    code_site_champ = champs.get("code_site_champ", "")
    if code_site_champ and code_site_champ != code_client:
        notes.append(
            f"Champ 'Code site' séparé ({code_site_champ!r}) différent du code client "
            f"de facturation ({code_client!r}) — DYNAMIZ_PHARMA.docx dit d'utiliser le "
            f"code client de facturation, vérifier qu'il n'y a pas d'incohérence."
        )

    enseigne = champs.get("nom_point_de_vente", "")
    ticket.customer.enseigne = enseigne
    ticket.customer.adresse = champs.get("adresse", "")

    code_postal, ville = extraire_cp_ville(champs.get("cp_ville", ""))
    ticket.customer.code_postal = code_postal
    ticket.customer.ville = ville

    ticket.customer.nom = champs.get("contact_site", "")
    ticket.customer.fixe = champs.get("telephone", "")

    ticket.intervention.type_intervention = "Contrat"
    if CONTRAT_DYNAMIZ:
        ticket.intervention.contrat = CONTRAT_DYNAMIZ
    else:
        notes.append(
            "Contrat DYNAMIZ PHARMA non renseigné automatiquement : un seul contrat "
            "existe côté Pivot, mais son libellé exact n'est pas documenté dans "
            "DYNAMIZ_PHARMA.docx — sélectionner manuellement, ou compléter la "
            "constante CONTRAT_DYNAMIZ dans cet agent."
        )

    type_demande = detecter_type_demande(champs.get("type_demande", ""))
    if not type_demande:
        notes.append(
            f"Type de demande non reconnu (valeur brute : {champs.get('type_demande', '')!r}, "
            f"attendu SAV/INSTALLATION/PREVISITE) — Type/Type de ticket à compléter manuellement."
        )
    ticket.intervention.type = type_demande
    ticket.intervention.type_ticket = "Incident" if type_demande == TYPE_MAINTENANCE_SAV else "Demande" if type_demande else ""

    description = champs.get("description", "")
    contexte, intervention_a_realiser = extraire_contexte_et_intervention(description)
    ticket.intervention.problematique = contexte
    if not contexte:
        notes.append("Section 'Contexte:' introuvable dans la Description — Problématique à compléter manuellement.")

    ticket.intervention.origine = "Email"
    ticket.intervention.numero_incident_client = champs.get("numero_incident_client", "")

    ticket.intervention.intitule = construire_intitule(type_demande, enseigne, ville)

    # --- Matériel ---
    besoin_materiel_brut = champs.get("besoin_materiel", "")
    ticket.logistics.besoin_materiel = _normaliser(besoin_materiel_brut) == "oui"
    if ticket.logistics.besoin_materiel:
        ticket.logistics.pieces = champs.get("reference_materiel", "")
        ticket.logistics.envoi_piece_par = "IRIS"
        ticket.logistics.consigne_livraison = "Pudo"
    materiel_a_recuperer = champs.get("materiel_a_recuperer", "")
    if materiel_a_recuperer:
        ticket.logistics.retour_piece = "Oui"
        ticket.logistics.commentaire_logistique = _ajouter_si_absent(
            ticket.logistics.commentaire_logistique, f"Matériel à récupérer : {materiel_a_recuperer}"
        )

    # --- Planification ---
    ticket.procedure.intervention_sur_site = True
    ticket.procedure.prise_rdv = False

    date_heure_brut = champs.get("date_heure", "")
    if date_heure_brut and ressemble_a_une_date(date_heure_brut):
        ticket.procedure.date_limite = date_heure_brut
    elif date_heure_brut:
        ticket.procedure.consignes_planification = _ajouter_si_absent(
            ticket.procedure.consignes_planification, date_heure_brut
        )
        notes.append(
            "'Date et heure intervention souhaitée' ne ressemble pas à une vraie date/"
            "heure (texte libre de consigne de contact) — classé en Consigne de "
            "planification plutôt qu'en Date limite."
        )

    contraintes = champs.get("contraintes", "")
    if contraintes:
        ticket.procedure.consignes_planification = _ajouter_si_absent(
            ticket.procedure.consignes_planification, f"Contraintes de mise en place : {contraintes}"
        )

    nombre_technicien_brut = champs.get("nombre_technicien", "")
    if nombre_technicien_brut.isdigit():
        ticket.procedure.nombre_techniciens = int(nombre_technicien_brut)
    else:
        notes.append("Nombre de technicien non reconnu comme un nombre — à vérifier manuellement.")

    duree_brut = champs.get("duree", "")
    if duree_brut:
        ticket.procedure.duree = f"{duree_brut}h"

    ticket.procedure.travail_attendu = intervention_a_realiser

    hauteur = champs.get("hauteur", "")
    outillage_parts = [hauteur] if hauteur else []
    if detecter_besoin_testeur_cable(description):
        outillage_parts.append("Testeur de câbles")
        notes.append("Testeur de câbles ajouté automatiquement (texte mentionnant la vérification d'un câble réseau actif/connecté).")
    ticket.procedure.autre_outillage = " ; ".join(outillage_parts)

    ticket.procedure.technicien_anglophone = False  # toujours Non (docx)
    ticket.procedure.procedure = False  # ⚠️ DYNAMIZ PHARMA = SEUL client avec Procédure=Non par défaut (cf. SKILL.md)

    documents = champs.get("documents", "")
    if documents:
        notes.append(f"Documents à joindre (mentionné dans le formulaire, sens non documenté pour ce client) : {documents!r}.")

    if notes:
        bloc_notes = "⚠️ Points à vérifier (générés automatiquement) :\n" + "\n".join(f"- {n}" for n in notes)
        ticket.intervention.commentaire_interne = _ajouter_si_absent(
            ticket.intervention.commentaire_interne, bloc_notes
        )

    return ticket