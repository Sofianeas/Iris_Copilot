"""
Agent ETAM.

Logique métier issue de ETAM.docx + TOKI_ETAM.txt.

⚠️ AVERTISSEMENT IMPORTANT : la structure de lecture (`lire_champs_etam`)
est déduite d'une CAPTURE D'ÉCRAN fournie par Sofiane, PAS d'un vrai fichier
.xlsx inspecté (contrairement à AXE E-SANTE où on a pu lire le vrai
fichier). La logique métier (branche Mobilité/Smartphone, règle A52,
LOG-SWAP-UPS...) est testée et fiable ; la lecture de cellules par
POSITION DE COLONNE (A=Label, B=Valeur, C=Valeur secondaire pour le seul
champ "Matériel à préparer") est une hypothèse à confirmer contre le vrai
fichier.

Couverture de cette version :
- Branche "Maintenance Mobilité" / sous-cas Smartphone Samsung (Swap
  Transporteur) : entièrement implémentée et testée contre l'exemple
  fourni (ETAM Lingerie 0351, Samsung A54 demandé -> A52 imposé).
- Branches TPE ETAM / IMAC / Slovaquie & RT : détection de branche
  implémentée, mais SEULES les règles communes sont appliquées.

Règles clés de la branche Mobilité/Smartphone :
- Type = "Swap Transporteur" -> pas d'intervention sur site.
- ⚠️ RÈGLE CRITIQUE (TOKI_ETAM.txt, "Consigne client") : prioriser
  SYSTÉMATIQUEMENT l'A52, MÊME SI Akkodis demande explicitement un A54/A56.
- "Panne" -> transporteur AFFRETEMENT SANS + référence LOG-SWAP-UPS.
  "Perte/Vol" -> transporteur CLIENT SANS, PAS de LOG-SWAP-UPS.
- Code client ETAM vs INVEST21 : INVEST21 si l'enseigne n'est PAS "Etam"
  elle-même.
- Technicien anglophone déduit du champ "Langue" du formulaire.
- "Noter le nom du matériel dans le champ TAG de FootPrints" (TOKI) :
  IGNORÉ -- FootPrints est l'ancien logiciel, non utilisé.

--- Intégration Rule Engine (cette étape) ---
  Ajout de 2 paramètres optionnels à `enrich_ticket_depuis_fichier`, en FIN
  de signature (après `texte_mail`) : `rag_decision: RagDecision | None =
  None` et `activer_rule_engine: bool = False`. Désactivé par défaut :
  rétrocompatibilité totale. ETAM n'a qu'UN SEUL point de sortie.
  ⚠️ Le contrat n'est JAMAIS vide sur cet agent (MOBILITE ->
  "Maintenance Mobilité", autres branches -> nom de la branche elle-même)
  -> `client_rule` (qui ne recommande que si contrat vide) ne se
  déclenchera jamais ici -- comportement attendu, testé explicitement.
"""

import re
import unicodedata

from openpyxl import load_workbook

from app.models.ticket import Ticket
from app.models.rag_decision import RagDecision
from app.services import rule_engine


def _sans_accents(texte: str) -> str:
    """Retire les accents pour fiabiliser les comparaisons de texte."""
    if not texte:
        return ""
    return "".join(
        c for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )


def _normaliser(texte: str) -> str:
    """Normalise (minuscule + sans accent + trim) pour des comparaisons robustes."""
    return _sans_accents((texte or "").strip().lower())


def _ajouter_si_absent(texte_existant: str, bloc: str) -> str:
    """Ajoute `bloc` à `texte_existant` s'il n'y est pas déjà (idempotence)."""
    if not bloc:
        return texte_existant
    if texte_existant and bloc in texte_existant:
        return texte_existant
    if texte_existant:
        return f"{texte_existant.strip()}\n\n{bloc}"
    return bloc


def _formater_recommandations_rule_engine(recommandations) -> str:
    """
    Formate les RuleRecommendation (rule_engine.executer) en un bloc de
    texte destiné à commentaire_interne -- ne modifie JAMAIS un champ
    métier directement (même politique que sur les autres agents migrés).
    """
    if not recommandations:
        return ""
    lignes = ["🧩 Recommandations du Rule Engine (à vérifier, jamais appliquées automatiquement) :"]
    for reco in recommandations:
        lignes.append(
            f"- Champ '{reco.field}' -> '{reco.value}' "
            f"(confiance={reco.confidence:.2f}, source={reco.source}) : {reco.reason}"
        )
    return "\n".join(lignes)


SECTIONS_CONNUES = {
    "information sur l'enseigne": "enseigne",
    "information technique du materiel en panne": "panne",
    "information technique du materiel a preparer": "preparer",
    "description de l'incident rencontre": "incident",
    "information n": "incident_akkodis",
}


def detecter_section(label_normalise: str) -> str:
    """Détecte une ligne de titre de section (cf. SECTIONS_CONNUES)."""
    for motif, cle in SECTIONS_CONNUES.items():
        if label_normalise.startswith(motif):
            return cle
    return ""


def construire_cle(section: str, label_normalise: str) -> str:
    """Mappe (section_courante, label) -> clé interne."""
    if label_normalise.startswith("code magasin"):
        return "code_magasin"
    if label_normalise.startswith("adresse de livraison"):
        return "adresse_livraison"
    if label_normalise.startswith("nom et n"):
        return "contact_nom"
    if label_normalise.startswith("adresse mail du site"):
        return "email_site"
    if label_normalise.startswith("materiel (modele)") or label_normalise.startswith("materiel(modele)"):
        if section == "panne":
            return "materiel_panne"
        if section == "preparer":
            return "materiel_preparer"
        return ""
    if label_normalise.startswith("numero de serie"):
        return "numero_serie"
    if label_normalise.startswith("pays"):
        return "pays"
    if label_normalise.startswith("langue"):
        return "langue"
    if label_normalise.startswith("typologie panne"):
        return "typologie_panne"
    if label_normalise.startswith("description panne"):
        return "description_panne"
    if label_normalise.startswith("tests effectues"):
        return "tests_effectues"
    if label_normalise.startswith("commentaires"):
        return "commentaires"
    if label_normalise.startswith("n") and "dossier akkodis" in label_normalise:
        return "numero_dossier_akkodis"
    return ""


def lire_champs_etam(fichier, feuille: str | None = None) -> dict:
    """Lit le fichier ETAM. Hypothèse de colonnes (NON VÉRIFIÉE contre un vrai fichier)."""
    wb = load_workbook(fichier, data_only=True)
    ws = wb[feuille] if feuille else wb.active

    resultat: dict[str, str] = {}
    section_courante = ""
    for row in ws.iter_rows():
        cells = list(row)
        label_brut = str(cells[0].value).strip() if len(cells) > 0 and cells[0].value else ""
        valeur_b = str(cells[1].value).strip() if len(cells) > 1 and cells[1].value else ""
        valeur_c = str(cells[2].value).strip() if len(cells) > 2 and cells[2].value else ""

        if not label_brut:
            continue
        label_normalise = _normaliser(label_brut)

        section_detectee = detecter_section(label_normalise)
        if section_detectee and not valeur_b:
            section_courante = section_detectee
            continue

        cle = construire_cle(section_courante, label_normalise)
        if not cle:
            continue

        if cle == "materiel_preparer":
            resultat["materiel_preparer_description"] = valeur_b
            resultat["materiel_preparer"] = valeur_c or valeur_b
        else:
            resultat[cle] = valeur_b

    return resultat


RE_ENSEIGNE_CODE = re.compile(r"^(.+?)\s*\(n°?\s*(\d+)\)", re.IGNORECASE)
RE_CP_VILLE = re.compile(r"^(\d{4,5})\s+(.+)$")
RE_TEL = re.compile(r"t[ée]l\s*:?\s*(.+)", re.IGNORECASE)
RE_PREFIXE_INTL_FR = re.compile(r"0033\s*\(0\)\s*")


def normaliser_telephone_fr_intl(numero: str) -> str:
    """Convertit '0033 (0) 4 67 64 04 46' -> '04 67 64 04 46' (format national FR)."""
    if not numero:
        return numero
    return RE_PREFIXE_INTL_FR.sub("0", numero).strip()


def parser_bloc_adresse_etam(bloc: str) -> dict:
    """Parse le bloc multi-lignes "Adresse de livraison et d'enlèvement"."""
    resultat = {
        "enseigne": "", "code_magasin": "", "complement_adresse": "",
        "adresse": "", "code_postal": "", "ville": "", "pays": "", "telephone_site": "",
    }
    if not bloc:
        return resultat

    lignes_restantes = []
    for ligne in (l.strip() for l in bloc.splitlines() if l.strip()):
        m_enseigne = RE_ENSEIGNE_CODE.match(ligne)
        if m_enseigne:
            resultat["enseigne"] = m_enseigne.group(1).strip()
            resultat["code_magasin"] = m_enseigne.group(2).strip()
            continue
        m_tel = RE_TEL.match(ligne)
        if m_tel:
            resultat["telephone_site"] = normaliser_telephone_fr_intl(m_tel.group(1).strip())
            continue
        m_cp_ville = RE_CP_VILLE.match(ligne)
        if m_cp_ville:
            resultat["code_postal"] = m_cp_ville.group(1).strip()
            resultat["ville"] = m_cp_ville.group(2).strip()
            continue
        lignes_restantes.append(ligne)

    if lignes_restantes:
        derniere = lignes_restantes[-1]
        if derniere.isupper() and len(derniere) < 30:
            resultat["pays"] = derniere
            lignes_restantes = lignes_restantes[:-1]
        if lignes_restantes:
            resultat["adresse"] = lignes_restantes[-1]
            if len(lignes_restantes) > 1:
                resultat["complement_adresse"] = " ".join(lignes_restantes[:-1])

    return resultat


BUNDLE_A52 = (
    "CETAMOB-SAMSUNG-A52-5G",
    "SMARTPHONE SAMSUNG A52-5G – ETAM – SPARE (inclut téléphone, verre trempé, coque, tour de cou)",
)
REF_LOG_SWAP_UPS = "LOG-SWAP-UPS"


def detecter_modele_samsung(texte: str) -> str:
    """A52 / A54 / A56, détecté dans le matériel demandé."""
    texte_n = _normaliser(texte)
    for modele in ("a52", "a54", "a56"):
        if modele in texte_n:
            return modele.upper()
    return ""


def determiner_transporteur_et_swap(typologie_panne: str) -> tuple[str, bool, str]:
    """Panne -> AFFRETEMENT SANS + LOG-SWAP-UPS. Perte/Vol -> CLIENT SANS, pas de swap."""
    texte = _normaliser(typologie_panne)
    if "perte" in texte or "vol" in texte:
        return "CLIENT SANS", False, (
            "Typologie 'Perte/Vol' détectée : transporteur 'CLIENT SANS' appliqué, "
            "PAS de LOG-SWAP-UPS demandé (cf. TOKI_ETAM.txt)."
        )
    return "AFFRETEMENT SANS", True, ""


def determiner_code_client(enseigne: str) -> tuple[str, bool]:
    """ETAM.docx : INVEST21 si l'enseigne n'est pas Etam elle-même."""
    if "etam" in _normaliser(enseigne):
        return "ETAM", True
    return "INVEST21", False


def technicien_anglophone_depuis_langue(langue: str) -> tuple[bool, bool]:
    """Retourne (technicien_anglophone, langue_reconnue)."""
    langue_n = _normaliser(langue)
    if not langue_n:
        return False, False
    if "francais" in langue_n or "french" in langue_n:
        return False, True
    return True, True


def detecter_branche(materiel_panne: str, materiel_preparer: str) -> tuple[str, bool]:
    """Déduit la branche ETAM (TPE / MOBILITE / IMAC / SLOVAQUIE_RT)."""
    texte = _normaliser(f"{materiel_panne} {materiel_preparer}")
    if any(mot in texte for mot in ("p400", "v400m", "adyen", "terminal de paiement", " tpe ", "tpe ")):
        return "TPE", True
    if any(mot in texte for mot in (
        "samsung", "smartphone", "ipad", "tablette", "douchette", "zebra",
        "pathfinder", "rfid", "sco", "kiosk", "tabletop",
    )):
        return "MOBILITE", True
    if any(mot in texte for mot in ("fermeture", "ouverture", "ajout materiel", "ramassage materiel", "transfert boutique")):
        return "IMAC", True
    return "MOBILITE", False


def enrich_ticket_depuis_fichier(
    ticket: Ticket,
    fichier_etam,
    texte_mail: str = "",
    rag_decision: RagDecision | None = None,
    activer_rule_engine: bool = False,
) -> Ticket:
    """
    Enrichit un Ticket à partir du fichier ETAM.

    `rag_decision` (optionnel) : une RagDecision déjà calculée en amont,
    transmise telle quelle au Rule Engine si celui-ci est activé.

    `activer_rule_engine` (par défaut False) : si True, exécute
    rule_engine.executer(ticket, rag_decision) et ajoute ses
    recommandations à commentaire_interne -- jamais à un champ métier.
    """
    notes: list[str] = []

    champs = lire_champs_etam(fichier_etam)

    bloc_adresse = parser_bloc_adresse_etam(champs.get("adresse_livraison", ""))

    code_magasin = champs.get("code_magasin") or bloc_adresse["code_magasin"]
    if champs.get("code_magasin") and bloc_adresse["code_magasin"] and champs["code_magasin"] != bloc_adresse["code_magasin"]:
        notes.append(
            f"Code magasin incohérent entre le champ dédié ({champs['code_magasin']!r}) "
            f"et le bloc adresse ({bloc_adresse['code_magasin']!r}) — à vérifier."
        )
    if code_magasin and not re.fullmatch(r"\d{4}", code_magasin.strip()):
        notes.append(f"Code magasin inhabituel ({code_magasin!r}) — attendu sur 4 chiffres, à vérifier.")
    ticket.customer.code_site = code_magasin

    enseigne = bloc_adresse["enseigne"]
    ticket.customer.enseigne = enseigne
    code_client, reconnu_etam = determiner_code_client(enseigne)
    ticket.customer.client = code_client
    if not reconnu_etam:
        notes.append(
            f"Enseigne '{enseigne}' différente de Etam -- code client 'INVEST21' "
            f"appliqué au lieu de 'ETAM' (cf. ETAM.docx). Vérifier qu'il ne s'agit "
            f"pas d'ERAM (client différent, à ne pas confondre)."
        )

    ticket.customer.adresse = bloc_adresse["adresse"]
    ticket.customer.complement_adresse = bloc_adresse["complement_adresse"]
    ticket.customer.code_postal = bloc_adresse["code_postal"]
    ticket.customer.ville = bloc_adresse["ville"]
    ticket.customer.pays = champs.get("pays") or bloc_adresse["pays"]
    ticket.customer.fixe = bloc_adresse["telephone_site"]

    nom_contact = champs.get("contact_nom", "")
    if nom_contact:
        ticket.customer.prenom = nom_contact
        notes.append(
            f"Contact '{nom_contact!r}' placé dans prenom (un seul nom donné dans "
            f"le formulaire, pas de nom de famille distinct) — à corriger si besoin."
        )
    else:
        notes.append("Nom du contact sur site introuvable — à compléter manuellement.")
    ticket.customer.email = champs.get("email_site", "")

    branche, branche_confiante = detecter_branche(champs.get("materiel_panne", ""), champs.get("materiel_preparer", ""))
    notes.append(
        f"⚠️ BRANCHE déduite : « {branche} »"
        + ("" if branche_confiante else " (PAR DÉFAUT, aucun mot-clé déterminant trouvé)")
        + " — À CONFIRMER avant saisie."
    )

    ticket.intervention.type_intervention = "Contrat"
    ticket.intervention.numero_serie = champs.get("numero_serie", "")
    ticket.intervention.problematique = champs.get("description_panne", "")
    ticket.intervention.origine = "Email"
    ticket.intervention.numero_incident_client = champs.get("numero_dossier_akkodis", "")
    if not ticket.intervention.numero_incident_client:
        notes.append("N° Dossier Akkodis introuvable — Numéro d'incident client à compléter manuellement.")

    langue_anglophone, langue_reconnue = technicien_anglophone_depuis_langue(champs.get("langue", ""))
    ticket.procedure.technicien_anglophone = langue_anglophone
    if not langue_reconnue:
        notes.append("Champ 'Langue' non renseigné/non reconnu — technicien_anglophone=Non appliqué par défaut, à vérifier.")

    travail_attendu_parties = []
    if champs.get("tests_effectues"):
        travail_attendu_parties.append(f"Tests déjà effectués sur site : {champs['tests_effectues']}")
    if champs.get("commentaires"):
        travail_attendu_parties.append(f"Commentaire du demandeur : {champs['commentaires']}")
    ticket.procedure.travail_attendu = "\n\n".join(travail_attendu_parties)

    if branche == "MOBILITE":
        modele_demande = detecter_modele_samsung(champs.get("materiel_preparer", "") + " " + champs.get("materiel_preparer_description", ""))

        ticket.intervention.contrat = "Maintenance Mobilité"
        ticket.intervention.type = "Swap Transporteur"
        ticket.procedure.intervention_sur_site = False

        if modele_demande in ("A54", "A56"):
            ticket.intervention.sous_type = "Smartphone"
            ticket.logistics.pieces = f"{BUNDLE_A52[0]} ({BUNDLE_A52[1]})"
            notes.append(
                f"⚠️ RÈGLE CLIENT APPLIQUÉE : le formulaire demande un {modele_demande}, "
                f"mais TOKI_ETAM.txt exige de prioriser SYSTÉMATIQUEMENT l'A52 sur "
                f"chaque ticket de maintenance, même si Akkodis demande un A54/A56. "
                f"Référence A52 substituée — sauf rupture de stock A52 confirmée (à vérifier)."
            )
        elif modele_demande == "A52":
            ticket.intervention.sous_type = "Smartphone"
            ticket.logistics.pieces = f"{BUNDLE_A52[0]} ({BUNDLE_A52[1]})"
        else:
            notes.append(
                "Matériel mobilité non reconnu comme un smartphone Samsung "
                "(A52/A54/A56) — Type/Sous-type et pièce à compléter manuellement."
            )

        ticket.logistics.besoin_materiel = True
        ticket.logistics.envoi_piece_par = "IRIS"
        transporteur, ajouter_swap, note_transporteur = determiner_transporteur_et_swap(champs.get("typologie_panne", ""))
        ticket.logistics.consigne_livraison = f"Transporteur : {transporteur}"
        if ajouter_swap and ticket.logistics.pieces:
            ticket.logistics.pieces = f"{ticket.logistics.pieces} + {REF_LOG_SWAP_UPS}"
        if note_transporteur:
            notes.append(note_transporteur)

        ville = ticket.customer.ville or ""
        ticket.intervention.intitule = " ".join(
            p for p in ("Maintenance Mobilité – Smartphone", enseigne, code_magasin, ville) if p
        )
        notes.append(
            "Intitulé construit sans règle stricte documentée pour la branche "
            "Mobilité — format best-effort, à ajuster si besoin."
        )

    else:
        notes.append(
            f"Branche '{branche}' détectée mais NON ENCORE IMPLÉMENTÉE en détail "
            f"dans cet agent (seule la branche Mobilité/Smartphone a été testée "
            f"contre un exemple réel à ce jour) — seules les règles communes "
            f"ont été appliquées. Fournir un exemple réel de cette branche pour la durcir."
        )
        ticket.intervention.contrat = branche

    if notes:
        bloc_notes = "⚠️ Points à vérifier (générés automatiquement) :\n" + "\n".join(f"- {n}" for n in notes)
        ticket.intervention.commentaire_interne = _ajouter_si_absent(
            ticket.intervention.commentaire_interne, bloc_notes
        )

    if activer_rule_engine:
        recommandations = rule_engine.executer(ticket, rag_decision)
        bloc_recommandations = _formater_recommandations_rule_engine(recommandations)
        ticket.intervention.commentaire_interne = _ajouter_si_absent(
            ticket.intervention.commentaire_interne, bloc_recommandations
        )

    return ticket

"""
Agent ETAM.

[... docstring métier inchangé ...]

--- Harmonisation Framework des Agents (P3-424) ---
Logique métier déplacée dans `_appliquer_regles_etam` (privée), appelée
par `EtamAgent.analyze()`. `analyze()` ne pilote jamais le Rule Engine
(Règle 9). `enrich_ticket_depuis_fichier()` reste un adaptateur de
compatibilité. `request.fichier_attache` porte directement le fichier
ETAM, même intégration native qu'AXE E-SANTE.
"""

import re
import unicodedata

from openpyxl import load_workbook

from app.agents.base_agent import BaseAgent
from app.models.agent_contracts import AgentRequest, AgentResult
from app.models.rag_decision import RagDecision
from app.models.ticket import Ticket
from app.services import rule_engine


def _sans_accents(texte: str) -> str:
    if not texte:
        return ""
    return "".join(
        c for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )


def _normaliser(texte: str) -> str:
    return _sans_accents((texte or "").strip().lower())


def _ajouter_si_absent(texte_existant: str, bloc: str) -> str:
    if not bloc:
        return texte_existant
    if texte_existant and bloc in texte_existant:
        return texte_existant
    if texte_existant:
        return f"{texte_existant.strip()}\n\n{bloc}"
    return bloc


def _formater_recommandations_rule_engine(recommandations) -> str:
    if not recommandations:
        return ""
    lignes = ["🧩 Recommandations du Rule Engine (à vérifier, jamais appliquées automatiquement) :"]
    for reco in recommandations:
        lignes.append(
            f"- Champ '{reco.field}' -> '{reco.value}' "
            f"(confiance={reco.confidence:.2f}, source={reco.source}) : {reco.reason}"
        )
    return "\n".join(lignes)


SECTIONS_CONNUES = {
    "information sur l'enseigne": "enseigne",
    "information technique du materiel en panne": "panne",
    "information technique du materiel a preparer": "preparer",
    "description de l'incident rencontre": "incident",
    "information n": "incident_akkodis",
}


def detecter_section(label_normalise: str) -> str:
    for motif, cle in SECTIONS_CONNUES.items():
        if label_normalise.startswith(motif):
            return cle
    return ""


def construire_cle(section: str, label_normalise: str) -> str:
    if label_normalise.startswith("code magasin"):
        return "code_magasin"
    if label_normalise.startswith("adresse de livraison"):
        return "adresse_livraison"
    if label_normalise.startswith("nom et n"):
        return "contact_nom"
    if label_normalise.startswith("adresse mail du site"):
        return "email_site"
    if label_normalise.startswith("materiel (modele)") or label_normalise.startswith("materiel(modele)"):
        if section == "panne":
            return "materiel_panne"
        if section == "preparer":
            return "materiel_preparer"
        return ""
    if label_normalise.startswith("numero de serie"):
        return "numero_serie"
    if label_normalise.startswith("pays"):
        return "pays"
    if label_normalise.startswith("langue"):
        return "langue"
    if label_normalise.startswith("typologie panne"):
        return "typologie_panne"
    if label_normalise.startswith("description panne"):
        return "description_panne"
    if label_normalise.startswith("tests effectues"):
        return "tests_effectues"
    if label_normalise.startswith("commentaires"):
        return "commentaires"
    if label_normalise.startswith("n") and "dossier akkodis" in label_normalise:
        return "numero_dossier_akkodis"
    return ""


def lire_champs_etam(fichier, feuille: str | None = None) -> dict:
    wb = load_workbook(fichier, data_only=True)
    ws = wb[feuille] if feuille else wb.active

    resultat: dict[str, str] = {}
    section_courante = ""
    for row in ws.iter_rows():
        cells = list(row)
        label_brut = str(cells[0].value).strip() if len(cells) > 0 and cells[0].value else ""
        valeur_b = str(cells[1].value).strip() if len(cells) > 1 and cells[1].value else ""
        valeur_c = str(cells[2].value).strip() if len(cells) > 2 and cells[2].value else ""

        if not label_brut:
            continue
        label_normalise = _normaliser(label_brut)

        section_detectee = detecter_section(label_normalise)
        if section_detectee and not valeur_b:
            section_courante = section_detectee
            continue

        cle = construire_cle(section_courante, label_normalise)
        if not cle:
            continue

        if cle == "materiel_preparer":
            resultat["materiel_preparer_description"] = valeur_b
            resultat["materiel_preparer"] = valeur_c or valeur_b
        else:
            resultat[cle] = valeur_b

    return resultat


RE_ENSEIGNE_CODE = re.compile(r"^(.+?)\s*\(n°?\s*(\d+)\)", re.IGNORECASE)
RE_CP_VILLE = re.compile(r"^(\d{4,5})\s+(.+)$")
RE_TEL = re.compile(r"t[ée]l\s*:?\s*(.+)", re.IGNORECASE)
RE_PREFIXE_INTL_FR = re.compile(r"0033\s*\(0\)\s*")


def normaliser_telephone_fr_intl(numero: str) -> str:
    if not numero:
        return numero
    return RE_PREFIXE_INTL_FR.sub("0", numero).strip()


def parser_bloc_adresse_etam(bloc: str) -> dict:
    resultat = {
        "enseigne": "", "code_magasin": "", "complement_adresse": "",
        "adresse": "", "code_postal": "", "ville": "", "pays": "", "telephone_site": "",
    }
    if not bloc:
        return resultat

    lignes_restantes = []
    for ligne in (l.strip() for l in bloc.splitlines() if l.strip()):
        m_enseigne = RE_ENSEIGNE_CODE.match(ligne)
        if m_enseigne:
            resultat["enseigne"] = m_enseigne.group(1).strip()
            resultat["code_magasin"] = m_enseigne.group(2).strip()
            continue
        m_tel = RE_TEL.match(ligne)
        if m_tel:
            resultat["telephone_site"] = normaliser_telephone_fr_intl(m_tel.group(1).strip())
            continue
        m_cp_ville = RE_CP_VILLE.match(ligne)
        if m_cp_ville:
            resultat["code_postal"] = m_cp_ville.group(1).strip()
            resultat["ville"] = m_cp_ville.group(2).strip()
            continue
        lignes_restantes.append(ligne)

    if lignes_restantes:
        derniere = lignes_restantes[-1]
        if derniere.isupper() and len(derniere) < 30:
            resultat["pays"] = derniere
            lignes_restantes = lignes_restantes[:-1]
        if lignes_restantes:
            resultat["adresse"] = lignes_restantes[-1]
            if len(lignes_restantes) > 1:
                resultat["complement_adresse"] = " ".join(lignes_restantes[:-1])

    return resultat


BUNDLE_A52 = (
    "CETAMOB-SAMSUNG-A52-5G",
    "SMARTPHONE SAMSUNG A52-5G – ETAM – SPARE (inclut téléphone, verre trempé, coque, tour de cou)",
)
REF_LOG_SWAP_UPS = "LOG-SWAP-UPS"


def detecter_modele_samsung(texte: str) -> str:
    texte_n = _normaliser(texte)
    for modele in ("a52", "a54", "a56"):
        if modele in texte_n:
            return modele.upper()
    return ""


def determiner_transporteur_et_swap(typologie_panne: str) -> tuple[str, bool, str]:
    texte = _normaliser(typologie_panne)
    if "perte" in texte or "vol" in texte:
        return "CLIENT SANS", False, (
            "Typologie 'Perte/Vol' détectée : transporteur 'CLIENT SANS' appliqué, "
            "PAS de LOG-SWAP-UPS demandé (cf. TOKI_ETAM.txt)."
        )
    return "AFFRETEMENT SANS", True, ""


def determiner_code_client(enseigne: str) -> tuple[str, bool]:
    if "etam" in _normaliser(enseigne):
        return "ETAM", True
    return "INVEST21", False


def technicien_anglophone_depuis_langue(langue: str) -> tuple[bool, bool]:
    langue_n = _normaliser(langue)
    if not langue_n:
        return False, False
    if "francais" in langue_n or "french" in langue_n:
        return False, True
    return True, True


def detecter_branche(materiel_panne: str, materiel_preparer: str) -> tuple[str, bool]:
    texte = _normaliser(f"{materiel_panne} {materiel_preparer}")
    if any(mot in texte for mot in ("p400", "v400m", "adyen", "terminal de paiement", " tpe ", "tpe ")):
        return "TPE", True
    if any(mot in texte for mot in (
        "samsung", "smartphone", "ipad", "tablette", "douchette", "zebra",
        "pathfinder", "rfid", "sco", "kiosk", "tabletop",
    )):
        return "MOBILITE", True
    if any(mot in texte for mot in ("fermeture", "ouverture", "ajout materiel", "ramassage materiel", "transfert boutique")):
        return "IMAC", True
    return "MOBILITE", False


def _appliquer_regles_etam(ticket: Ticket, fichier_etam, texte_mail: str = "") -> Ticket:
    """Logique métier ETAM pure (extraite de l'ancien `enrich_ticket_depuis_fichier`, sans le bloc Rule Engine)."""
    notes: list[str] = []

    champs = lire_champs_etam(fichier_etam)

    bloc_adresse = parser_bloc_adresse_etam(champs.get("adresse_livraison", ""))

    code_magasin = champs.get("code_magasin") or bloc_adresse["code_magasin"]
    if champs.get("code_magasin") and bloc_adresse["code_magasin"] and champs["code_magasin"] != bloc_adresse["code_magasin"]:
        notes.append(
            f"Code magasin incohérent entre le champ dédié ({champs['code_magasin']!r}) "
            f"et le bloc adresse ({bloc_adresse['code_magasin']!r}) — à vérifier."
        )
    if code_magasin and not re.fullmatch(r"\d{4}", code_magasin.strip()):
        notes.append(f"Code magasin inhabituel ({code_magasin!r}) — attendu sur 4 chiffres, à vérifier.")
    ticket.customer.code_site = code_magasin

    enseigne = bloc_adresse["enseigne"]
    ticket.customer.enseigne = enseigne
    code_client, reconnu_etam = determiner_code_client(enseigne)
    ticket.customer.client = code_client
    if not reconnu_etam:
        notes.append(
            f"Enseigne '{enseigne}' différente de Etam -- code client 'INVEST21' "
            f"appliqué au lieu de 'ETAM' (cf. ETAM.docx). Vérifier qu'il ne s'agit "
            f"pas d'ERAM (client différent, à ne pas confondre)."
        )

    ticket.customer.adresse = bloc_adresse["adresse"]
    ticket.customer.complement_adresse = bloc_adresse["complement_adresse"]
    ticket.customer.code_postal = bloc_adresse["code_postal"]
    ticket.customer.ville = bloc_adresse["ville"]
    ticket.customer.pays = champs.get("pays") or bloc_adresse["pays"]
    ticket.customer.fixe = bloc_adresse["telephone_site"]

    nom_contact = champs.get("contact_nom", "")
    if nom_contact:
        ticket.customer.prenom = nom_contact
        notes.append(
            f"Contact '{nom_contact!r}' placé dans prenom (un seul nom donné dans "
            f"le formulaire, pas de nom de famille distinct) — à corriger si besoin."
        )
    else:
        notes.append("Nom du contact sur site introuvable — à compléter manuellement.")
    ticket.customer.email = champs.get("email_site", "")

    branche, branche_confiante = detecter_branche(champs.get("materiel_panne", ""), champs.get("materiel_preparer", ""))
    notes.append(
        f"⚠️ BRANCHE déduite : « {branche} »"
        + ("" if branche_confiante else " (PAR DÉFAUT, aucun mot-clé déterminant trouvé)")
        + " — À CONFIRMER avant saisie."
    )

    ticket.intervention.type_intervention = "Contrat"
    ticket.intervention.numero_serie = champs.get("numero_serie", "")
    ticket.intervention.problematique = champs.get("description_panne", "")
    ticket.intervention.origine = "Email"
    ticket.intervention.numero_incident_client = champs.get("numero_dossier_akkodis", "")
    if not ticket.intervention.numero_incident_client:
        notes.append("N° Dossier Akkodis introuvable — Numéro d'incident client à compléter manuellement.")

    langue_anglophone, langue_reconnue = technicien_anglophone_depuis_langue(champs.get("langue", ""))
    ticket.procedure.technicien_anglophone = langue_anglophone
    if not langue_reconnue:
        notes.append("Champ 'Langue' non renseigné/non reconnu — technicien_anglophone=Non appliqué par défaut, à vérifier.")

    travail_attendu_parties = []
    if champs.get("tests_effectues"):
        travail_attendu_parties.append(f"Tests déjà effectués sur site : {champs['tests_effectues']}")
    if champs.get("commentaires"):
        travail_attendu_parties.append(f"Commentaire du demandeur : {champs['commentaires']}")
    ticket.procedure.travail_attendu = "\n\n".join(travail_attendu_parties)

    if branche == "MOBILITE":
        modele_demande = detecter_modele_samsung(champs.get("materiel_preparer", "") + " " + champs.get("materiel_preparer_description", ""))

        ticket.intervention.contrat = "Maintenance Mobilité"
        ticket.intervention.type = "Swap Transporteur"
        ticket.procedure.intervention_sur_site = False

        if modele_demande in ("A54", "A56"):
            ticket.intervention.sous_type = "Smartphone"
            ticket.logistics.pieces = f"{BUNDLE_A52[0]} ({BUNDLE_A52[1]})"
            notes.append(
                f"⚠️ RÈGLE CLIENT APPLIQUÉE : le formulaire demande un {modele_demande}, "
                f"mais TOKI_ETAM.txt exige de prioriser SYSTÉMATIQUEMENT l'A52 sur "
                f"chaque ticket de maintenance, même si Akkodis demande un A54/A56. "
                f"Référence A52 substituée — sauf rupture de stock A52 confirmée (à vérifier)."
            )
        elif modele_demande == "A52":
            ticket.intervention.sous_type = "Smartphone"
            ticket.logistics.pieces = f"{BUNDLE_A52[0]} ({BUNDLE_A52[1]})"
        else:
            notes.append(
                "Matériel mobilité non reconnu comme un smartphone Samsung "
                "(A52/A54/A56) — Type/Sous-type et pièce à compléter manuellement."
            )

        ticket.logistics.besoin_materiel = True
        ticket.logistics.envoi_piece_par = "IRIS"
        transporteur, ajouter_swap, note_transporteur = determiner_transporteur_et_swap(champs.get("typologie_panne", ""))
        ticket.logistics.consigne_livraison = f"Transporteur : {transporteur}"
        if ajouter_swap and ticket.logistics.pieces:
            ticket.logistics.pieces = f"{ticket.logistics.pieces} + {REF_LOG_SWAP_UPS}"
        if note_transporteur:
            notes.append(note_transporteur)

        ville = ticket.customer.ville or ""
        ticket.intervention.intitule = " ".join(
            p for p in ("Maintenance Mobilité – Smartphone", enseigne, code_magasin, ville) if p
        )
        notes.append(
            "Intitulé construit sans règle stricte documentée pour la branche "
            "Mobilité — format best-effort, à ajuster si besoin."
        )

    else:
        notes.append(
            f"Branche '{branche}' détectée mais NON ENCORE IMPLÉMENTÉE en détail "
            f"dans cet agent (seule la branche Mobilité/Smartphone a été testée "
            f"contre un exemple réel à ce jour) — seules les règles communes "
            f"ont été appliquées. Fournir un exemple réel de cette branche pour la durcir."
        )
        ticket.intervention.contrat = branche

    if notes:
        bloc_notes = "⚠️ Points à vérifier (générés automatiquement) :\n" + "\n".join(f"- {n}" for n in notes)
        ticket.intervention.commentaire_interne = _ajouter_si_absent(
            ticket.intervention.commentaire_interne, bloc_notes
        )

    return ticket


class EtamAgent(BaseAgent):
    """Agent ETAM conforme au contrat BaseAgent. `request.fichier_attache` porte directement le fichier ETAM."""

    def analyze(self, request: AgentRequest) -> AgentResult:
        try:
            ticket = _appliquer_regles_etam(
                request.ticket, request.fichier_attache, texte_mail=request.texte_mail
            )
            return AgentResult(ticket=ticket, succes=True)
        except Exception as exc:
            return AgentResult(ticket=request.ticket, succes=False, erreur=str(exc))


_AGENT = EtamAgent()


def enrich_ticket_depuis_fichier(
    ticket: Ticket,
    fichier_etam,
    texte_mail: str = "",
    rag_decision: RagDecision | None = None,
    activer_rule_engine: bool = False,
) -> Ticket:
    """⚠️ ADAPTATEUR DE COMPATIBILITÉ -- délègue à EtamAgent.analyze(), reproduit ici l'ancien comportement de activer_rule_engine."""
    request = AgentRequest(ticket=ticket, texte_mail=texte_mail, fichier_attache=fichier_etam, rag_decision=rag_decision)
    result = _AGENT.analyze(request)
    ticket_resultat = result.ticket

    if activer_rule_engine:
        recommandations = rule_engine.executer(ticket_resultat, rag_decision)
        bloc_recommandations = _formater_recommandations_rule_engine(recommandations)
        ticket_resultat.intervention.commentaire_interne = _ajouter_si_absent(
            ticket_resultat.intervention.commentaire_interne, bloc_recommandations
        )
    return ticket_resultat

